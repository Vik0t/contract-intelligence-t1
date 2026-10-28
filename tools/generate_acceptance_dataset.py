"""Generate fictional contracts, gold and a split before pipeline tuning (TZ 3.2)."""
from __future__ import annotations

import argparse
from datetime import date
from hashlib import sha256
import json
from pathlib import Path
import random

from docx import Document
import pymupdf


def generate(destination: Path, *, count=36, seed=20261027):
    if count < 20:
        raise ValueError('ТЗ требует не менее 20 договоров')
    destination.mkdir(parents=True, exist_ok=False)
    rng = random.Random(seed)
    manifest = {'version': 1, 'seed': seed, 'generator_sha256': sha256(Path(__file__).read_bytes()).hexdigest(),
                'gold_review': 'pending_manual_review', 'documents': []}
    font = pymupdf.Font('cjk')
    for index in range(count):
        ident = index + 1
        template = index % 3
        split = 'holdout' if (index // 3) % 3 == 2 else 'development'
        atypical = template == 2
        number = f'{rng.randint(100, 999)}/{ident:03d}-Т'
        day = date(2026, rng.randint(1, 9), rng.randint(1, 28))
        roles = [('Заказчик', 'Исполнитель'), ('Покупатель', 'Поставщик'), ('Арендатор', 'Арендодатель')][template]
        customer = f'ООО "Модельный заказчик {ident:03d}"'
        contractor = f'ИП Модельный Исполнитель {ident:03d}'
        amount = rng.randint(10, 800) * 1000 + rng.randint(1, 99) / 100
        currency = ['RUB', 'USD', None][index % 3]
        money = f'{amount:,.2f}'.replace(',', ' ').replace('.', ',')
        currency_text = {'RUB': 'рублей', 'USD': 'долларов США', None: ''}[currency]
        heading = f'ДОГОВОР № {number}'
        date_line = f'г. Новосибирск {day:%d.%m.%Y}'
        party_lines = [f'{roles[0]}: {customer}', f'{roles[1]}: {contractor}']
        total = f'Общая сумма договора составляет {money} {currency_text}.'.strip()
        distractors = ['Счёт № 999 от 01.01.2020.', 'Аванс 5 000 рублей. Штраф 300 рублей. НДС 20%.',
                       'Цена за единицу 100 рублей.', 'Представитель действует по доверенности № 55 от 02.02.2020.']
        lines = [heading, date_line, *party_lines, 'Предмет договора: модельные услуги.', total]
        if atypical:
            lines = [heading, date_line, *distractors, 'РЕКВИЗИТЫ СТОРОН', *party_lines, total]
        doc = Document()
        if template == 1:
            doc.sections[0].header.paragraphs[0].text = heading
            doc.add_paragraph(date_line)
            table = doc.add_table(rows=2, cols=2)
            for row, role, party in zip(table.rows, roles, [customer, contractor]):
                row.cells[0].text, row.cells[1].text = role, party
            doc.add_paragraph('Предмет договора: модельные услуги.')
            doc.add_paragraph(total)
        else:
            for line in lines:
                doc.add_paragraph(line)
        filenames = {ext: f'contract_{ident:03d}.{ext}' for ext in ('pdf', 'docx')}
        doc.save(destination / filenames['docx'])
        with pymupdf.open() as pdf:
            page = pdf.new_page()
            page.insert_font(fontname='fictional', fontbuffer=font.buffer)
            for row, line in enumerate(lines):
                page.insert_text((50, 65 + row * 27), line, fontsize=10, fontname='fictional')
            pdf.subset_fonts()
            pdf.save(destination / filenames['pdf'], garbage=4, deflate=True)
        manifest['documents'].append({'id': ident, 'split': split, 'kind': 'atypical' if atypical else 'typical',
            'template_group': f'model-{template + 1}', 'files': filenames,
            'sha256': {ext: sha256((destination / name).read_bytes()).hexdigest() for ext, name in filenames.items()},
            'gold': {'contract_number': number, 'contract_date': day.isoformat(),
                     'parties': [{'name': customer, 'role': roles[0], 'canonical_role': 'Заказчик'},
                                 {'name': contractor, 'role': roles[1], 'canonical_role': 'Исполнитель'}],
                     'total_amount': {'value': amount, 'currency': currency}}})
    payload = json.dumps(manifest, ensure_ascii=False, indent=2)
    (destination / 'manifest.json').write_text(payload, encoding='utf-8')
    (destination / 'manifest.sha256').write_text(sha256(payload.encode()).hexdigest() + '\n')
    return {'contracts': count, 'templates': 3, 'development': sum(d['split'] == 'development' for d in manifest['documents']),
            'holdout': sum(d['split'] == 'holdout' for d in manifest['documents']), 'manifest_sha256': sha256(payload.encode()).hexdigest()}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--count', type=int, default=36)
    parser.add_argument('--seed', type=int, default=20261027)
    args = parser.parse_args()
    print(json.dumps(generate(args.output, count=args.count, seed=args.seed), ensure_ascii=False))
