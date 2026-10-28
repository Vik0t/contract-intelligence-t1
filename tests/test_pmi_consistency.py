from datetime import date

from docx import Document
import pymupdf
import pytest

from backend.file_validation import UploadValidationError, validate_document_file
from src.extraction.reliability import assess_fields
from src.extraction.validation import validate_model_result, merge_extractions
from src.parsing import parse_document, document_ir_to_text


@pytest.mark.parametrize('value', [0, -100])
def test_nonpositive_contract_total_is_never_found(value):
    result = validate_model_result({'total_amount': {'value': value, 'evidence': f'Сумма договора {value} руб.'}}, f'Сумма договора {value} руб.')
    assert result['total_amount']['status'] == 'needs_review'


def test_date_is_checked_against_upload_day_not_reprocessing_day():
    result = {'contract_date': {'value': '2026-10-02', 'evidence': '02.10.2026'}}
    assess_fields(result, reference_date=date(2026, 10, 1))
    assert result['contract_date']['status'] == 'needs_review'


def test_identical_parties_are_not_found_and_original_roles_are_kept():
    parties = [{'name': 'ООО Альфа', 'role': role, 'name_evidence': 'ООО Альфа'} for role in ['Покупатель', 'Поставщик']]
    result = assess_fields({'parties': parties})
    assert result['_quality']['parties']['status'] == 'needs_review'
    assert parties[0]['role'] == 'Покупатель'
    assert parties[0]['canonical_role'] == 'Заказчик'


def test_two_totals_on_one_page_keep_both_sources():
    text = 'PAGE 1\nОбщая сумма договора составляет 1 250 000 руб.\nОбщая сумма договора составляет 1 500 000 руб.'
    result = validate_model_result({'total_amount': {'value': 1250000, 'currency': 'RUB', 'evidence': 'Общая сумма договора составляет 1 250 000 руб.'}}, text)
    assert result['total_amount']['status'] == 'needs_review'
    assert {c['value'] for c in result['total_amount']['candidates']} == {1250000, 1500000}


def test_conflicting_chunks_retain_sources_and_do_not_claim_found():
    results = [validate_model_result({'total_amount': {'value': value, 'evidence': f'Цена договора {value} рублей'}}, f'PAGE {i}\nЦена договора {value} рублей') for i, value in enumerate([10, 20], 1)]
    merged = merge_extractions(results)
    assert merged['total_amount']['status'] == 'needs_review'
    assert {c['page_number'] for c in merged['total_amount']['candidates']} == {1, 2}


def test_currency_is_not_inferred_when_not_explicit():
    result = validate_model_result({'total_amount': {'value': 100, 'currency': 'RUB', 'evidence': 'Сумма договора 100'}}, 'Сумма договора 100')
    assert result['total_amount']['currency'] is None
    assert result['total_amount']['status'] == 'found'


def test_real_citation_does_not_ground_a_different_model_value():
    result = validate_model_result({'total_amount': {'value': 100000, 'currency': 'RUB', 'evidence': 'Сумма договора 1000 рублей'}}, 'Сумма договора 1000 рублей')
    assert result['total_amount']['status'] == 'needs_review'


def test_docx_headers_tables_footers_and_paragraph_locations(tmp_path):
    path = tmp_path / 'contract.docx'
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = 'Договор № HEADER-7'
    doc.add_paragraph('Наиме-\nнование «организации»')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Заказчик ООО Альфа'
    doc.sections[0].footer.paragraphs[0].text = 'Исполнитель ООО Бета'
    doc.save(path)
    parsed = parse_document(path)
    text = document_ir_to_text(parsed)
    assert text.index('HEADER-7') < text.index('Наименование') < text.index('Заказчик') < text.index('Исполнитель')
    assert '«' not in text
    result = assess_fields({'contract_number': {'value': 'HEADER-7', 'evidence': 'Договор № HEADER-7'}}, document=parsed)
    assert result['contract_number']['paragraph_number'] == 1


@pytest.mark.parametrize('encrypted', [False, True])
def test_no_text_and_password_protected_pdf_are_rejected(tmp_path, encrypted):
    path = tmp_path / 'bad.pdf'
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        if encrypted:
            page.insert_text((72, 72), 'Secret contract')
        pdf.save(path, **({'encryption': pymupdf.PDF_ENCRYPT_AES_256, 'user_pw': 'secret', 'owner_pw': 'owner'} if encrypted else {}))
    with pytest.raises(UploadValidationError, match='паролем|извлекаемого текста'):
        validate_document_file(path, '.pdf', max_pdf_pages=20)


def test_empty_docx_is_rejected(tmp_path):
    path = tmp_path / 'empty.docx'
    Document().save(path)
    with pytest.raises(UploadValidationError, match='извлекаемого текста'):
        validate_document_file(path, '.docx', max_pdf_pages=20)
