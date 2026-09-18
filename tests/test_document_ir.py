import pytest
from docx import Document
from PIL import Image
from pydantic import ValidationError

from src.document_ir import (
    BlockIR,
    BlockKind,
    BlockSource,
    DocumentIR,
    PageIR,
    SourceFormat,
)
from src.local_extraction import validate_extraction
from src.parsing import document_ir_to_text, parse_document


def test_page_count_must_match_pages() -> None:
    with pytest.raises(ValidationError):
        DocumentIR(
            source_format="pdf",
            source_path="contract.pdf",
            page_count=1,
            pages=[],
        )


def test_docx_is_parsed_to_document_ir(tmp_path) -> None:
    path = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph("Договор поставки № 15")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Поставщик"
    table.cell(0, 1).text = "ООО Альфа"
    document.save(path)

    parsed = parse_document(path)

    assert parsed.source_format == SourceFormat.DOCX
    assert parsed.page_count == 1
    assert [block.kind for block in parsed.pages[0].blocks] == [
        BlockKind.PARAGRAPH,
        BlockKind.TABLE,
    ]
    assert "Поставщик | ООО Альфа" in parsed.pages[0].blocks[1].text


def test_photo_is_parsed_with_ocr_and_orientation(tmp_path) -> None:
    path = tmp_path / "contract.jpg"
    Image.new("RGB", (900, 1200), "white").save(path, format="JPEG")

    class FakeOCR:
        def extract_blocks(self, image, *, page_number, page_width, page_height):
            assert image.size == (900, 1200)
            return [
                BlockIR(
                    id="photo-ocr",
                    kind=BlockKind.PARAGRAPH,
                    text="Договор поставки № 42",
                    source=BlockSource.OCR,
                    order=0,
                )
            ]

    parsed = parse_document(path, ocr_backend=FakeOCR())

    assert parsed.source_format == SourceFormat.IMAGE
    assert parsed.page_count == 1
    assert parsed.pages[0].needs_ocr is True
    assert document_ir_to_text(parsed) == "PAGE 1\n\nДоговор поставки № 42"


def test_document_text_prefers_ocr_on_scanned_page() -> None:
    document = DocumentIR(
        source_format=SourceFormat.PDF,
        source_path="contract.pdf",
        page_count=1,
        pages=[
            PageIR(
                number=1,
                width=500,
                height=700,
                needs_ocr=True,
                blocks=[
                    BlockIR(
                        id="digital",
                        kind=BlockKind.PARAGRAPH,
                        text="короткий цифровой слой",
                        source=BlockSource.DIGITAL,
                    ),
                    BlockIR(
                        id="ocr",
                        kind=BlockKind.PARAGRAPH,
                        text="полный распознанный договор",
                        source=BlockSource.OCR,
                        order=0,
                    ),
                ],
            )
        ],
    )

    text = document_ir_to_text(document)

    assert text == "PAGE 1\n\nполный распознанный договор"


def test_evidence_repairs_date_and_page_number() -> None:
    document = DocumentIR(
        source_format=SourceFormat.PDF,
        source_path="contract.pdf",
        page_count=2,
        pages=[
            PageIR(
                number=1,
                blocks=[
                    BlockIR(
                        id="page-1",
                        kind=BlockKind.PARAGRAPH,
                        text="Договор заключён 23.03.2026.",
                        source=BlockSource.DIGITAL,
                    )
                ],
            ),
            PageIR(
                number=2,
                blocks=[
                    BlockIR(
                        id="page-2",
                        kind=BlockKind.PARAGRAPH,
                        text="Договор действует до 20.08.2026.",
                        source=BlockSource.DIGITAL,
                    )
                ],
            ),
        ],
    )
    result = {
        "contract_date": {
            "value": "2024-03-23",
            "evidence": "23.03.2026",
            "page_number": 2,
        },
        "term": {
            "value": "до 20.08.2026",
            "evidence": "Договор действует до 20.08.2026.",
            "page_number": 1,
        },
    }

    validated = validate_extraction(result, document)

    assert validated["contract_date"]["value"] == "2026-03-23"
    assert validated["contract_date"]["page_number"] == 1
    assert validated["term"]["page_number"] == 2
