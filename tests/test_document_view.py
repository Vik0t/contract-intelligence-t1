from pathlib import Path
import shutil

import pymupdf
import pytest
from docx import Document
from PIL import Image

from backend.document_view import locate_evidence, page_count, render_page


def test_pdf_renders_original_page_and_locates_exact_fragment(tmp_path):
    path = tmp_path / "source.pdf"
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 100), "First page")
    pdf.new_page().insert_text((72, 350), "Unique contract clause")
    pdf.save(path)
    pdf.close()

    assert page_count(path) == 2
    location = locate_evidence(path, "Unique contract clause", preferred_page=1)
    assert location["page"] == 2
    assert 0.3 < location["rects"][0]["y"] < 0.5
    assert render_page(path, 2).startswith(b"\x89PNG")


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice unavailable")
def test_docx_uses_real_rendered_pages(tmp_path):
    path = tmp_path / "formatted.docx"
    document = Document()
    document.add_paragraph("First contract page")
    document.add_page_break()
    document.add_paragraph("Unique second page clause")
    document.save(path)

    assert page_count(path) == 2
    location = locate_evidence(path, "Unique second page clause", preferred_page=1)
    assert location["page"] == 2
    assert Path(str(path) + ".rendered.pdf").is_file()


def test_unicode_pdf_hyphens_do_not_break_real_source_highlight(tmp_path):
    from scripts.ml_check import fixtures
    source = fixtures(tmp_path)["pdf"]
    location = locate_evidence(source, "ДОГОВОР ПОСТАВКИ № НГУ-2026-017", 1)
    assert location["page"] == 1
    assert location["rects"][0]["width"] > 0


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice unavailable")
def test_docx_cyrillic_font_runs_do_not_break_number_highlight(tmp_path):
    path = tmp_path / 'font-runs.docx'
    doc = Document()
    doc.add_paragraph('ДОГОВОР № 596/007-Т')
    doc.add_paragraph('Заказчик: ООО «Модельный заказчик 007»')
    doc.save(path)
    location = locate_evidence(path, 'ДОГОВОР № 596/007-Т', preferred_page=1)
    assert location is not None
    assert location['page'] == 1
    assert locate_evidence(path, 'ООО "Модельный заказчик 007"', preferred_page=1) is not None


def test_photo_highlight_reuses_original_ocr_coordinates(tmp_path):
    path = tmp_path / "photo.jpg"
    Image.new("RGB", (1000, 1200), "white").save(path)
    ir = {"pages": [{"number": 1, "width": 1000, "height": 1200, "blocks": [{
        "text": "Договор № 17", "source": "ocr", "bbox": {"x0": 100, "y0": 200, "x1": 600, "y1": 240},
    }]}]}
    assert page_count(path) == 1
    assert render_page(path, 1).startswith(b"\x89PNG")
    location = locate_evidence(path, "Договор № 17", 1, document_ir=ir)
    assert location["rects"][0]["x"] == pytest.approx(.1)
    assert location["rects"][0]["y"] == pytest.approx(200 / 1200)
    assert locate_evidence(path, "Несуществующая цитата", document_ir=ir) is None
