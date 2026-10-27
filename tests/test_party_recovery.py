import json
from copy import deepcopy
from types import SimpleNamespace

from src.document_ir import DocumentIR, PageIR
from src.extraction.party_resolution import merge_parties, link_from_source, source_links
from src.extraction.validation import ContractResult, validate_model_result, merge_extractions
from src.extraction.recovery import recover_missing


def test_full_and_short_legal_names_merge_without_detaching_inn():
    result = merge_parties([{'name':'Общество с ограниченной ответственностью «Ромашка»','inn':None}],
                           [{'name':'ООО "Ромашка"','inn':'1234567890'}])
    assert len(result)==1 and result[0]['inn']=='1234567890'


def test_same_inn_does_not_collapse_different_named_companies():
    rows=[{'name':name,'inn':'1234567890'} for name in ('ООО «Ромашка»','АО «Ромашка»')]
    assert len(merge_parties([], rows))==2
    assert link_from_source(rows, '')[1]


def test_two_inns_for_one_company_are_not_silently_resolved():
    rows=[{'name':'ООО «Ромашка»','inn':inn} for inn in ('1234567890','9876543210')]
    assert len(merge_parties([],rows))==2
    assert link_from_source(rows,'')[1]


def test_anonymous_inn_is_attached_only_to_explicit_name_record():
    rows=[{'name':'ООО «Ромашка»','inn':None},{'name':None,'inn':'1234567890','inn_evidence':'ИНН: 1234567890'}]
    text='PAGE 8\nПОСТАВЩИК:\nООО «Ромашка»\n\nИНН: 1234567890'
    result,issues=link_from_source(rows,text)
    assert len(result)==1 and result[0]['inn']=='1234567890' and not issues
    assert result[0]['name']=='ООО «Ромашка»'


def test_flattened_two_column_names_cannot_be_linked_by_nearest_inn():
    rows=[{'name':name,'inn':None} for name in ('Alpha LLC','Beta LLC')]
    text='PAGE 2\nAlpha LLC Beta LLC\nИНН: 1234567890 ИНН: 9876543210'
    assert not source_links(rows,text)
    assert all(p['inn'] is None for p in link_from_source(rows,text)[0])


def test_pdf_column_geometry_establishes_identity_independent_of_output_order(tmp_path,monkeypatch):
    import pymupdf
    path=tmp_path/'source.pdf';path.write_bytes(b'synthetic-layout-fixture')
    def word(x,y,text):return (x,y,x+len(text)*5,y+10,text,0,0,0)
    words=[word(10,10,'Реквизиты сторон'),word(10,30,'Alpha LLC'),word(250,30,'Beta LLC'),
           word(10,50,'ИНН:'),word(40,50,'1234567890'),word(250,50,'ИНН:'),word(280,50,'9876543210')]
    class PDF:
        def __enter__(self):return [SimpleNamespace(get_text=lambda *a,**kw:words)]
        def __exit__(self,*args):pass
    monkeypatch.setattr(pymupdf,'open',lambda *a,**kw:PDF())
    document=DocumentIR(source_format='pdf',source_path=str(path),pages=[PageIR(number=1)])
    text='PAGE 1\nРеквизиты сторон\nAlpha LLC Beta LLC\nИНН: 1234567890 ИНН: 9876543210'
    rows=[{'name':name,'inn':None} for name in ('Beta LLC','Alpha LLC')]
    result,issues=link_from_source(rows,text,document)
    assert {p['name']:p['inn'] for p in result}=={'Alpha LLC':'1234567890','Beta LLC':'9876543210'}
    assert all(p['inn_link']['method']=='same_requisites_column' for p in result) and not issues


def test_native_docx_records_survive_rendered_column_interleaving(tmp_path):
    from docx import Document
    path=tmp_path/'source.docx';doc=Document()
    for name,inn in [('Alpha LLC','1234567890'),('Beta LLC','9876543210')]:
        doc.add_paragraph(name);doc.add_paragraph('ИНН: '+inn)
    doc.save(path)
    text='PAGE 1\nAlpha LLC Beta LLC\nИНН: 1234567890 ИНН: 9876543210'
    ir=DocumentIR(source_format='docx',source_path=str(path),pages=[PageIR(number=1)])
    result,_=link_from_source([{'name':n,'inn':None} for n in ('Beta LLC','Alpha LLC')],text,ir)
    assert {p['name']:p['inn'] for p in result}=={'Alpha LLC':'1234567890','Beta LLC':'9876543210'}


def test_final_merge_runs_source_normalization_on_full_document():
    result=merge_extractions([validate_model_result({'term':{'value':None}},'PAGE 1\nДоговор')],
                             'PAGE 1\nДоговор действует до\n14 июля 2033 года.')
    assert result['term']['value']=='2033-07-14'


def test_repair_is_bounded_and_scalar_values_remain_unchanged():
    text='PAGE 1\nМесто подписания: г. Самара\nДата подписания: 14.07.2033\nДоговор № Т-19\nPAGE 2\nОбщая сумма договора 700 рублей.'
    first=validate_model_result({'total_amount':{'value':700,'evidence':'Общая сумма договора 700 рублей.'}},text)
    previous=deepcopy(first);calls=[]
    def request(prompt):
        calls.append(prompt)
        return json.dumps({'contract_date':{'value':'2033-07-14','evidence':'Дата подписания: 14.07.2033'},
                           'place':{'value':'Самара','evidence':'Место подписания: г. Самара'},
                           'total_amount':{'value':999,'evidence':'Общая сумма договора 700 рублей.'}})
    result,trace=recover_missing(first,text,request,ContractResult.model_json_schema())
    assert len(calls)==trace['calls']==1 and result['contract_date']['value']=='2033-07-14'
    assert result['place']['value']=='Самара' and result['total_amount']['value']==700
    assert first==previous


def test_repair_rejects_wrong_date_value_and_passport_source():
    text='PAGE 1\nДата подписания: 14.07.2033\nДоговор № Т-19\nПаспорт выдан 01.01.2020.'
    first=validate_model_result({'contract_date':{'value':None}},text)
    for evidence in ('Дата подписания: 14.07.2033','Паспорт выдан 01.01.2020.'):
        result,trace=recover_missing(first,text,lambda _:json.dumps({'contract_date':{'value':'2020-01-01','evidence':evidence}}),ContractResult.model_json_schema())
        assert result['contract_date']['value']=='2033-07-14'
        assert 'contract_date' not in trace['accepted']
        assert result['contract_date']['method']=='rule'


def test_repair_never_assumes_two_names_and_two_inns_have_matching_order():
    text='PAGE 1\nСТОРОНЫ: Alpha LLC Beta LLC\nИНН: 1234567890 ИНН: 9876543210'
    first=validate_model_result({'parties':[]},text)
    raw={'parties':[{'name':n,'inn':inn,'name_evidence':n,'inn_evidence':'ИНН: '+inn} for n,inn in [('Beta LLC','1234567890'),('Alpha LLC','9876543210')]]}
    result,_=recover_missing(first,text,lambda _:json.dumps(raw),ContractResult.model_json_schema())
    assert len(result['parties'])==2
    assert all(p.get('inn') is None for p in result['parties'])


def test_repair_failure_preserves_the_first_pass(monkeypatch):
    from src.extraction import chat_pipeline
    monkeypatch.setenv('EXTRACTION_API_MODEL','synthetic-model')
    monkeypatch.setenv('EXTRACTION_API_REPAIR_ENABLED','true')
    calls=[]
    def generate(*args,**kwargs):
        calls.append(1)
        if len(calls)>1:
            raise ConnectionError('Synthetic unavailable repair service')
        return json.dumps({'contract_number':{'value':'Т-19','evidence':'Договор № Т-19'}})
    monkeypatch.setattr(chat_pipeline,'generate_chat',generate)
    result,_=chat_pipeline.extract_fields_with_api('PAGE 1\nДоговор № Т-19\nСТОРОНЫ: Alpha LLC Beta LLC\nИНН: 1234567890 ИНН: 9876543210')
    assert result['contract_number']['value']=='Т-19'
    assert result['contract_date']['value'] is None
    assert result['_repair']['status']=='error' and len(calls)==2
